"""只读加载、校验并插值题目附件1与附件2。"""

from __future__ import annotations

from bisect import bisect_right
from dataclasses import dataclass
from hashlib import sha256
from math import isclose, isfinite
from pathlib import Path
from typing import Sequence

from openpyxl import load_workbook


class ModelInputError(ValueError):
    """附件结构、单位或数值不符合 V4 约定。"""


@dataclass(frozen=True)
class SourceTrace:
    filename: str
    worksheet: str
    sha256: str
    byte_size: int
    records: int


def default_attachment_directory() -> Path:
    """Locate the sibling competition attachments without embedding a machine path."""

    return Path(__file__).resolve().parents[2].parent / "A题" / "附件"


def _file_trace(path: Path, records: int) -> SourceTrace:
    digest = sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return SourceTrace(path.name, "Sheet1", digest.hexdigest(), path.stat().st_size, records)


def _numeric_rows(
    path: Path, columns: int, expected_headers: tuple[str, ...]
) -> tuple[tuple[str, ...], tuple[tuple[float, ...], ...]]:
    if not path.is_file():
        raise ModelInputError(f"找不到输入文件：{path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["Sheet1"]:
            raise ModelInputError(f"{path.name} 必须且只能包含工作表 Sheet1。")
        sheet = workbook["Sheet1"]
        if sheet.max_column != columns:
            raise ModelInputError(f"{path.name} 应有 {columns} 列，实际为 {sheet.max_column} 列。")
        headers = tuple(str(sheet.cell(1, column).value or "").strip() for column in range(1, columns + 1))
        if headers != expected_headers:
            raise ModelInputError(
                f"{path.name} 表头必须为 {expected_headers}，实际为 {headers}。"
            )
        rows: list[tuple[float, ...]] = []
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if len(values) != columns or any(value is None for value in values):
                raise ModelInputError(f"{path.name} 第 {row_number} 行存在缺失值。")
            try:
                numeric = tuple(float(value) for value in values)
            except (TypeError, ValueError) as exc:
                raise ModelInputError(f"{path.name} 第 {row_number} 行含非数值。") from exc
            if any(not isfinite(value) for value in numeric):
                raise ModelInputError(f"{path.name} 第 {row_number} 行含非有限数。")
            rows.append(numeric)
        return headers, tuple(rows)
    finally:
        workbook.close()


def _validate_time_axis(
    times: Sequence[float], *, records: int, step_s: float, end_s: float, filename: str
) -> None:
    if len(times) != records:
        raise ModelInputError(f"{filename} 应有 {records} 条记录，实际为 {len(times)} 条。")
    if not isclose(times[0], 0.0, abs_tol=1.0e-9) or not isclose(times[-1], end_s, abs_tol=1.0e-9):
        raise ModelInputError(f"{filename} 时间范围必须为 0 至 {end_s:g} s。")
    for index, (left, right) in enumerate(zip(times, times[1:]), start=2):
        if right <= left:
            raise ModelInputError(f"{filename} 时间在第 {index} 条数据处不严格递增。")
        if not isclose(right - left, step_s, rel_tol=0.0, abs_tol=1.0e-9):
            raise ModelInputError(f"{filename} 时间步长必须恒为 {step_s:g} s。")


def _linear_value(times: Sequence[float], values: Sequence[float], time_s: float) -> float:
    position = bisect_right(times, time_s)
    if position == 0:
        return values[0]
    if position >= len(times):
        return values[-1]
    left = position - 1
    fraction = (time_s - times[left]) / (times[position] - times[left])
    return values[left] + fraction * (values[position] - values[left])


@dataclass(frozen=True)
class EnvironmentInput:
    times_s: tuple[float, ...]
    temperatures_c: tuple[float, ...]
    equilibrium_moistures: tuple[float, ...]
    trace: SourceTrace
    nominal_temperature_c: float = 50.0
    nominal_equilibrium_moisture: float = 0.05

    def __post_init__(self) -> None:
        _validate_time_axis(
            self.times_s, records=241, step_s=60.0, end_s=14_400.0,
            filename=self.trace.filename,
        )
        if not (len(self.times_s) == len(self.temperatures_c) == len(self.equilibrium_moistures)):
            raise ModelInputError("附件1三列长度不一致。")
        if any(not isfinite(value) for value in (*self.temperatures_c, *self.equilibrium_moistures)):
            raise ModelInputError("附件1环境数据必须全部为有限数。")
        if any(value <= -273.15 for value in self.temperatures_c):
            raise ModelInputError("附件1温度不得低于绝对零度。")
        if any(value < 0.0 for value in self.equilibrium_moistures):
            raise ModelInputError("附件1空气水分数值不得为负。")
        if not isfinite(self.nominal_temperature_c) or self.nominal_temperature_c <= -273.15:
            raise ModelInputError("附件1尾部平台温度必须有限且高于绝对零度。")
        if not isfinite(self.nominal_equilibrium_moisture) or self.nominal_equilibrium_moisture < 0.0:
            raise ModelInputError("附件1尾部平台平衡含水率必须为有限非负数。")

    def at(self, time_s: float) -> tuple[float, float]:
        """V4 main policy: piecewise linear through 4 h, then (50, 0.05)."""

        if not isfinite(time_s) or time_s < 0.0:
            raise ValueError("查询时间必须是有限非负数。")
        if time_s > self.times_s[-1]:
            return self.nominal_temperature_c, self.nominal_equilibrium_moisture
        return (
            _linear_value(self.times_s, self.temperatures_c, time_s),
            _linear_value(self.times_s, self.equilibrium_moistures, time_s),
        )

    def last_hour_mean(self) -> tuple[float, float]:
        first = bisect_right(self.times_s, self.times_s[-1] - 3600.0 - 1.0e-9)
        temperatures = self.temperatures_c[first:]
        moistures = self.equilibrium_moistures[first:]
        return sum(temperatures) / len(temperatures), sum(moistures) / len(moistures)


@dataclass(frozen=True)
class RadiusInput:
    times_s: tuple[float, ...]
    radii_m: tuple[float, ...]
    trace: SourceTrace

    def __post_init__(self) -> None:
        _validate_time_axis(
            self.times_s, records=145, step_s=1800.0, end_s=259_200.0,
            filename=self.trace.filename,
        )
        if len(self.times_s) != len(self.radii_m):
            raise ModelInputError("附件2两列长度不一致。")
        if any(not isfinite(radius) for radius in self.radii_m):
            raise ModelInputError("附件2半径必须全部为有限数。")
        if any(radius <= 0.0 for radius in self.radii_m):
            raise ModelInputError("附件2半径必须为正。")
        if any(right > left + 1.0e-12 for left, right in zip(self.radii_m, self.radii_m[1:])):
            raise ModelInputError("附件2半径不得随时间增大。")
        if not isclose(self.radii_m[0], 0.02, abs_tol=1.0e-12):
            raise ModelInputError("附件2初始半径必须为 2.000 cm。")
        if not isclose(self.radii_m[-1], 0.01198, abs_tol=1.0e-12):
            raise ModelInputError("附件2末半径必须为 1.198 cm。")

    def at(self, time_s: float) -> float:
        """V4 main policy: piecewise linear through 72 h, then freeze at 1.198 cm."""

        if not isfinite(time_s) or time_s < 0.0:
            raise ValueError("查询时间必须是有限非负数。")
        if time_s > self.times_s[-1]:
            return self.radii_m[-1]
        return _linear_value(self.times_s, self.radii_m, time_s)


def load_environment(path: Path | None = None) -> EnvironmentInput:
    source = Path(path) if path is not None else default_attachment_directory() / "附件1.xlsx"
    _, rows = _numeric_rows(source, columns=3, expected_headers=("时间", "温度", "水分浓度"))
    times, temperatures, moistures = zip(*rows)
    return EnvironmentInput(tuple(times), tuple(temperatures), tuple(moistures), _file_trace(source, len(rows)))


def load_radius(path: Path | None = None) -> RadiusInput:
    source = Path(path) if path is not None else default_attachment_directory() / "附件2.xlsx"
    _, rows = _numeric_rows(source, columns=2, expected_headers=("时间", "半径"))
    times, radii_cm = zip(*rows)
    radii_m = tuple(value * 1.0e-2 for value in radii_cm)
    return RadiusInput(tuple(times), radii_m, _file_trace(source, len(rows)))
