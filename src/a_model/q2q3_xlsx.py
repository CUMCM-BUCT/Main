"""Validated, memory-bounded conversion of formal Q2/Q3 CSVs to XLSX."""

from copy import copy
import csv
from dataclasses import dataclass
import hashlib
import json
from math import isclose, isfinite
import os
from pathlib import Path
import tempfile

import xlsxwriter

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter


CSV_NAMES = ("result2_temperature.csv", "result2_moisture.csv", "result3_moisture.csv")
CSV_HEADER = ("time_s", *(f"r_{i / 10:g}_cm" for i in range(21)))
XLSX_HEADER = ("时间\\到药材中心的距离", *(i / 10 for i in range(21)))
MAX_XLSX_ROWS = 1_048_576
NUMBER_FORMAT = "0.0000"


@dataclass(frozen=True)
class CsvStats:
    rows: int
    first: tuple[float, ...]
    last: tuple[float, ...]
    sha256: str


def _numeric_row(row, path, line):
    if len(row) != len(CSV_HEADER):
        raise ValueError(f"{path.name}:{line} has {len(row)} columns; expected {len(CSV_HEADER)}.")
    try:
        values = tuple(float(value) for value in row)
    except ValueError as exc:
        raise ValueError(f"{path.name}:{line} contains a non-numeric value.") from exc
    if not all(isfinite(value) for value in values):
        raise ValueError(f"{path.name}:{line} contains a non-finite value.")
    return values


def _digest(path):
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def _scan(path, cadence_s):
    first = last = None
    rows = 0
    terminal_seen = False
    with path.open("r", encoding="utf-8", newline="") as stream:
        reader = csv.reader(stream)
        header = tuple(next(reader, ()))
        if header != CSV_HEADER:
            raise ValueError(f"{path.name} has an unexpected header: {header!r}.")
        for line, row in enumerate(reader, 2):
            if terminal_seen:
                raise ValueError(f"{path.name} has data after its off-cadence terminal row.")
            values = _numeric_row(row, path, line)
            rows += 1
            expected = rows * cadence_s
            if values[0] != expected:
                lower = expected - cadence_s
                if not lower < values[0] < expected:
                    raise ValueError(f"{path.name}:{line} violates the {cadence_s:g} s output cadence.")
                terminal_seen = True
            if last is not None and values[0] <= last[0]:
                raise ValueError(f"{path.name}:{line} time is not strictly increasing.")
            first = first or values
            last = values
    if rows == 0:
        raise ValueError(f"{path.name} has no data rows.")
    if rows + 1 > MAX_XLSX_ROWS:
        raise ValueError(
            f"{path.name} needs {rows + 1:,} Excel rows; the worksheet limit is {MAX_XLSX_ROWS:,}."
        )
    return CsvStats(rows, first, last, _digest(path))


def inspect_sources(source_dir):
    source_dir = Path(source_dir)
    paths = {name: source_dir / name for name in CSV_NAMES}
    missing = [name for name, path in paths.items() if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"Missing Q2/Q3 CSVs: {', '.join(missing)}")
    stats = {
        CSV_NAMES[0]: _scan(paths[CSV_NAMES[0]], 1),
        CSV_NAMES[1]: _scan(paths[CSV_NAMES[1]], 1),
        CSV_NAMES[2]: _scan(paths[CSV_NAMES[2]], 60),
    }
    temperature, moisture, q3 = (stats[name] for name in CSV_NAMES)
    if temperature.rows != moisture.rows or temperature.first[0] != moisture.first[0] or temperature.last[0] != moisture.last[0]:
        raise ValueError("The two result2 CSVs do not share one time grid.")
    return stats


def validate_formal_metadata(source_dir, stats):
    path = Path(source_dir) / "metadata.json"
    metadata = json.loads(path.read_text(encoding="utf-8"))
    if metadata.get("result_classification") != "formal":
        raise ValueError("Metadata does not classify this export as formal.")
    if metadata.get("numerical_settings", {}).get("convergence_verified") is not True:
        raise ValueError("Metadata does not record a passed convergence verification.")
    if metadata.get("source_files_unchanged") is not True:
        raise ValueError("Metadata does not confirm unchanged source files.")
    result = metadata.get("result", {})
    if result.get("status") != "event":
        raise ValueError("Metadata does not record a completed Q3 event.")
    event_time = result.get("event", {}).get("state", {}).get("time_s")
    if event_time != stats[CSV_NAMES[2]].last[0]:
        raise ValueError("Metadata event time does not match the result3 terminal row.")
    q2_scope = metadata.get("q2_scope_end_s")
    if q2_scope is not None and q2_scope != stats[CSV_NAMES[0]].last[0]:
        raise ValueError("Metadata Q2 scope does not match the result2 terminal row.")
    recorded = metadata.get("checkpoint", {}).get("files", {})
    for name, stat in stats.items():
        if recorded.get(name, {}).get("rows") != stat.rows or recorded.get(name, {}).get("sha256") != stat.sha256:
            raise ValueError(f"Metadata provenance does not match {name}.")
    return metadata


def _copy_style(source, target, number_format=None):
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)
    target.number_format = number_format or source.number_format


def _styled_cell(sheet, value, source, number_format=None):
    cell = WriteOnlyCell(sheet, value=value)
    _copy_style(source, cell, number_format)
    return cell


def _excel_number(value):
    return int(value) if value.is_integer() else value


def _write_book(template_path, sheet_sources, destination):
    template = load_workbook(template_path, data_only=False, read_only=False)
    book = Workbook(write_only=True)
    book.properties = copy(template.properties)
    try:
        for sheet_name, csv_path in sheet_sources:
            if sheet_name not in template.sheetnames:
                raise ValueError(f"Template is missing sheet {sheet_name!r}.")
            source_sheet = template[sheet_name]
            sheet = book.create_sheet(sheet_name)
            sheet.sheet_properties = copy(source_sheet.sheet_properties)
            sheet.sheet_state = source_sheet.sheet_state
            sheet.page_margins = copy(source_sheet.page_margins)
            sheet.page_setup = copy(source_sheet.page_setup)
            sheet.print_options = copy(source_sheet.print_options)
            sheet.freeze_panes = source_sheet.freeze_panes
            sheet.sheet_view.showGridLines = source_sheet.sheet_view.showGridLines
            sheet.sheet_format.defaultRowHeight = source_sheet.row_dimensions[1].height or source_sheet.sheet_format.defaultRowHeight
            sheet.column_dimensions["A"].width = source_sheet.column_dimensions["A"].width
            numeric_width = source_sheet.column_dimensions["B"].width
            for column in range(2, len(XLSX_HEADER) + 1):
                sheet.column_dimensions[get_column_letter(column)].width = numeric_width
            header = [_styled_cell(sheet, value, source_sheet.cell(1, 1 if column == 1 else 2))
                      for column, value in enumerate(XLSX_HEADER, 1)]
            sheet.append(header)
            with Path(csv_path).open("r", encoding="utf-8", newline="") as stream:
                reader = csv.reader(stream)
                next(reader)
                for line, row in enumerate(reader, 2):
                    values = _numeric_row(row, Path(csv_path), line)
                    cells = [_styled_cell(
                        sheet, _excel_number(value), source_sheet.cell(2, 1 if column == 1 else 2), NUMBER_FORMAT
                    ) for column, value in enumerate(values, 1)]
                    sheet.append(cells)
        book.save(destination)
    finally:
        template.close()


def verify_workbook(path, expected):
    book = load_workbook(path, data_only=False, read_only=True)
    try:
        if book.sheetnames != list(expected):
            raise ValueError(f"Unexpected sheet order in {Path(path).name}: {book.sheetnames!r}.")
        for sheet_name, stat in expected.items():
            sheet = book[sheet_name]
            rows = sheet.iter_rows()
            header = tuple(cell.value for cell in next(rows))
            if header != XLSX_HEADER:
                raise ValueError(f"Unexpected header in {Path(path).name}/{sheet_name}.")
            count = 0
            first = last = None
            for row in rows:
                count += 1
                values = tuple(cell.value for cell in row)
                if len(values) != len(XLSX_HEADER) or not all(type(value) in (int, float) and isfinite(value) for value in values):
                    raise ValueError(f"Non-numeric data in {Path(path).name}/{sheet_name}, row {count + 1}.")
                if any(cell.data_type in ("e", "f") for cell in row):
                    raise ValueError(f"Formula or Excel error in {Path(path).name}/{sheet_name}, row {count + 1}.")
                if any(cell.number_format != NUMBER_FORMAT for cell in row):
                    raise ValueError(f"Wrong number format in {Path(path).name}/{sheet_name}, row {count + 1}.")
                first = first or tuple(float(value) for value in values)
                last = tuple(float(value) for value in values)
            endpoints_match = all(
                isclose(actual, source, rel_tol=1e-14, abs_tol=1e-14)
                for actual, source in zip((*first, *last), (*stat.first, *stat.last))
            )
            if count != stat.rows or not endpoints_match:
                raise ValueError(f"Row count or endpoint values differ in {Path(path).name}/{sheet_name}.")
    finally:
        book.close()


def export_workbooks(source_dir, result2_template, result3_template, output_dir):
    source_dir, output_dir = Path(source_dir), Path(output_dir)
    stats = inspect_sources(source_dir)
    validate_formal_metadata(source_dir, stats)
    output_dir.mkdir(parents=True, exist_ok=True)
    result2, result3 = output_dir / "result2.xlsx", output_dir / "result3.xlsx"
    if result2.exists() or result3.exists():
        raise FileExistsError("result2.xlsx or result3.xlsx already exists; choose a new output directory.")
    temporaries = []
    try:
        for template, sheets, target in (
            (result2_template, (("温度", source_dir / CSV_NAMES[0]), ("水分浓度", source_dir / CSV_NAMES[1])), result2),
            (result3_template, (("Sheet1", source_dir / CSV_NAMES[2]),), result3),
        ):
            handle = tempfile.NamedTemporaryFile(prefix=f".{target.stem}-", suffix=".xlsx", dir=output_dir, delete=False)
            handle.close()
            temporary = Path(handle.name)
            temporaries.append(temporary)
            _write_book(template, sheets, temporary)
        verify_workbook(temporaries[0], {"温度": stats[CSV_NAMES[0]], "水分浓度": stats[CSV_NAMES[1]]})
        verify_workbook(temporaries[1], {"Sheet1": stats[CSV_NAMES[2]]})
        temporaries[0].rename(result2)
        temporaries[1].rename(result3)
        return result2, result3
    finally:
        for path in temporaries:
            if path.exists():
                path.unlink()


def export_workbooks_fast(source_dir, result2_template, result3_template, output_dir):
    """Write the validated contract with xlsxwriter's constant-memory path."""
    source_dir, output_dir = Path(source_dir), Path(output_dir)
    stats = inspect_sources(source_dir)
    validate_formal_metadata(source_dir, stats)
    output_dir.mkdir(parents=True, exist_ok=True)
    result2, result3 = output_dir / "result2.xlsx", output_dir / "result3.xlsx"
    if result2.exists() or result3.exists():
        raise FileExistsError("result2.xlsx or result3.xlsx already exists; choose a new output directory.")
    template2 = load_workbook(result2_template, read_only=False)
    template3 = load_workbook(result3_template, read_only=False)
    try:
        if tuple(template2.sheetnames) != ("温度", "水分浓度") or tuple(template3.sheetnames) != ("Sheet1",):
            raise ValueError("Templates do not have the required worksheet names.")
        width_a = template2["温度"].column_dimensions["A"].width or 19.625
        width_b = template2["温度"].column_dimensions["B"].width or 9.625
        width_3a = template3["Sheet1"].column_dimensions["A"].width or width_a
        width_3b = template3["Sheet1"].column_dimensions["B"].width or width_b
    finally:
        template2.close()
        template3.close()

    header = list(XLSX_HEADER)
    def write_book(path, sheets, widths):
        book = xlsxwriter.Workbook(path, {"constant_memory": True})
        header_format = book.add_format({"bold": True, "align": "center", "border": 1})
        number_format = book.add_format({"num_format": NUMBER_FORMAT})
        try:
            for sheet_name, csv_name in sheets:
                sheet = book.add_worksheet(sheet_name)
                sheet.set_column(0, 0, widths[0])
                sheet.set_column(1, len(XLSX_HEADER) - 1, widths[1])
                sheet.write_row(0, 0, header, header_format)
                with (source_dir / csv_name).open("r", encoding="utf-8", newline="") as stream:
                    reader = csv.reader(stream)
                    next(reader)
                    for row_index, row in enumerate(reader, 1):
                        values = _numeric_row(row, source_dir / csv_name, row_index + 1)
                        sheet.write_number(row_index, 0, _excel_number(values[0]), number_format)
                        for column, value in enumerate(values[1:], 1):
                            sheet.write_number(row_index, column, _excel_number(value), number_format)
        finally:
            book.close()

    write_book(result2, (("温度", CSV_NAMES[0]), ("水分浓度", CSV_NAMES[1])), (width_a, width_b))
    write_book(result3, (("Sheet1", CSV_NAMES[2]),), (width_3a, width_3b))
    verify_workbook(result2, {"温度": stats[CSV_NAMES[0]], "水分浓度": stats[CSV_NAMES[1]]})
    verify_workbook(result3, {"Sheet1": stats[CSV_NAMES[2]]})
    return result2, result3
