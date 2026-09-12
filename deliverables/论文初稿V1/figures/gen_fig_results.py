"""Generate publication figures for the four-question drying-model paper.

The script reads the exported solver tables rather than re-running the model.
All figures are written as vector PDF and 300-dpi PNG for convenient reuse.
"""
from pathlib import Path
import csv
import json
import math

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
MAIN = HERE.parents[2]
ATTACH = MAIN.parent / "A题" / "附件"
OUT = HERE

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "axes.unicode_minus": False,
    "font.size": 9,
    "axes.titlesize": 10,
    "axes.labelsize": 9,
    "legend.fontsize": 8,
    "xtick.labelsize": 8,
    "ytick.labelsize": 8,
    "figure.dpi": 120,
    "savefig.bbox": "tight",
})

COLORS = ["#0072B2", "#D55E00", "#009E73", "#CC79A7", "#E69F00"]
STYLES = ["-", "--", "-.", ":", (0, (3, 1, 1, 1))]


def savefig(name: str):
    fig = plt.gcf()
    fig.savefig(OUT / f"{name}.pdf", format="pdf")
    fig.savefig(OUT / f"{name}.png", format="png", dpi=300)
    plt.close(fig)


def read_csv(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def f(row, key):
    return float(row[key])


def col_at(row, name):
    return f(row, name)


def nearest(rows, time_s):
    return min(rows, key=lambda r: abs(f(r, "time_s") - time_s))


def plot_environment():
    wb = load_workbook(ATTACH / "附件1.xlsx", data_only=True, read_only=True)
    ws = wb.active
    vals = list(ws.iter_rows(min_row=2, values_only=True))
    t = np.array([float(x[0]) / 3600 for x in vals])
    tinf = np.array([float(x[1]) for x in vals])
    ce = np.array([float(x[2]) for x in vals])
    fig, ax = plt.subplots(1, 2, figsize=(6.7, 2.65), constrained_layout=True)
    ax[0].plot(t, tinf, color=COLORS[0], lw=1.8)
    ax[0].axhline(50, color="#666666", lw=1, ls=":", label="post-data extension")
    ax[0].set(xlabel="time / h", ylabel=r"$T_\infty$ / $^\circ$C", title="(a) Air temperature")
    ax[0].legend(frameon=False, loc="lower right")
    ax[1].plot(t, ce, color=COLORS[1], lw=1.8)
    ax[1].axhline(0.05, color="#666666", lw=1, ls=":", label="post-data extension")
    ax[1].set(xlabel="time / h", ylabel=r"$C_e$ / kg kg$^{-1}$", title="(b) Equivalent equilibrium moisture")
    ax[1].legend(frameon=False, loc="lower right")
    for a in ax:
        a.grid(alpha=0.25)
        a.spines[["top", "right"]].set_visible(False)
    savefig("fig01_environment")


def profile_plot(rows, times, temp_prefix, moist_prefix, name, xlim=(0, 2), time_scale=1.0):
    r_cm = np.arange(0, 2.01, 0.1)
    fig, ax = plt.subplots(1, 2, figsize=(6.7, 2.75), constrained_layout=True)
    for j, t in enumerate(times):
        row = nearest(rows, t * time_scale)
        temp = np.array([f(row, f"{temp_prefix}{x:g}_cm_C") for x in r_cm])
        moist = np.array([f(row, f"{moist_prefix}{x:g}_cm") for x in r_cm])
        lab = f"{t:g} h" if time_scale == 3600 else f"{t:g} s"
        ax[0].plot(r_cm, temp, color=COLORS[j], ls=STYLES[j], lw=1.55, label=lab)
        ax[1].plot(r_cm, moist, color=COLORS[j], ls=STYLES[j], lw=1.55, label=lab)
    ax[0].set(xlim=xlim, xlabel="radius / cm", ylabel=r"$T$ / $^\circ$C", title="(a) Temperature profiles")
    ax[1].set(xlim=xlim, xlabel="radius / cm", ylabel=r"$C$ / kg kg$^{-1}$", title="(b) Moisture profiles")
    for a in ax:
        a.grid(alpha=0.25)
        a.spines[["top", "right"]].set_visible(False)
    ax[0].legend(frameon=False, ncol=2)
    ax[1].legend(frameon=False, ncol=2)
    savefig(name)


def plot_q1():
    rows_t = read_csv(MAIN / "data/processed/q1/phase4_verified/q1_formal_temperature.csv")
    rows_c = read_csv(MAIN / "data/processed/q1/phase4_verified/q1_formal_moisture.csv")
    r_cm = np.arange(0, 2.01, 0.1)
    times = [100, 600, 1200, 1800]
    fig, ax = plt.subplots(1, 2, figsize=(6.7, 2.75), constrained_layout=True)
    for j, t in enumerate(times):
        rt = nearest(rows_t, t)
        rc = nearest(rows_c, t)
        temp = np.array([f(rt, f"r_{x:g}_cm") for x in r_cm])
        moist = np.array([f(rc, f"r_{x:g}_cm") for x in r_cm])
        lab = f"{t:g} s"
        ax[0].plot(r_cm, temp, color=COLORS[j], ls=STYLES[j], lw=1.55, label=lab)
        ax[1].plot(r_cm, moist, color=COLORS[j], ls=STYLES[j], lw=1.55, label=lab)
    ax[0].set(xlabel="radius / cm", ylabel=r"$T$ / $^\circ$C", title="(a) Temperature profiles")
    ax[1].set(xlabel="radius / cm", ylabel=r"$C$ / kg kg$^{-1}$", title="(b) Moisture profiles")
    for a in ax:
        a.grid(alpha=0.25)
        a.spines[["top", "right"]].set_visible(False)
    ax[0].legend(frameon=False, ncol=2)
    ax[1].legend(frameon=False, ncol=2)
    savefig("fig02_q1_profiles")


def plot_q2():
    rows = read_csv(MAIN / "data/processed/q2q3/convergence_graded_v2/N320_dt3p75_grade2/samples.csv")
    profile_plot(rows, [0.5, 1.0, 2.0, 3.0], "temperature_r_", "moisture_r_", "fig03_q2_profiles", time_scale=3600)


def plot_q3():
    rows = read_csv(MAIN / "data/processed/q2q3/convergence_graded_v2/N320_dt3p75_grade2/samples.csv")
    t = np.array([f(r, "time_s") / 3600 for r in rows])
    center = np.array([f(r, "moisture_center") for r in rows])
    surface = np.array([f(r, "moisture_surface") for r in rows])
    mean = np.array([f(r, "moisture_mean") for r in rows])
    fig, ax = plt.subplots(figsize=(6.7, 2.85), constrained_layout=True)
    ax.plot(t, center, color=COLORS[0], lw=1.8, label="center (controlling point)")
    ax.plot(t, mean, color=COLORS[2], lw=1.6, ls="--", label="cross-sectional mean")
    ax.plot(t, surface, color=COLORS[1], lw=1.6, ls="-.", label="surface")
    event_h = 206914.89990234375 / 3600
    ax.axhline(0.15, color="#444444", lw=1.1, ls=":", label="criterion $C=0.15$")
    ax.axvline(event_h, color="#666666", lw=1.0, ls=":")
    ax.scatter([event_h], [0.15], color=COLORS[0], s=24, zorder=4)
    ax.annotate(f"{event_h:.3f} h", xy=(event_h, 0.15), xytext=(-45, 18), textcoords="offset points",
                arrowprops=dict(arrowstyle="->", lw=0.8), fontsize=8)
    ax.set(xlabel="time / h", ylabel=r"moisture $C$ / kg kg$^{-1}$", title="Full-domain drying event (Q3)")
    ax.set_xlim(0, 60)
    ax.set_ylim(0, 2.7)
    ax.grid(alpha=0.25)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False, ncol=2, loc="upper right")
    savefig("fig04_q3_drying")


def get_event(n, dt):
    dt_tag = str(int(dt)) if float(dt).is_integer() else str(dt).replace('.', 'p')
    p = MAIN / "data/processed/q2q3/convergence_graded_v2" / f"N{n}_dt{dt_tag}_grade2" / "summary.json"
    with p.open("r", encoding="utf-8") as fobj:
        j = json.load(fobj)
    return j["result"]["event"]["state"]["time_s"] / 3600


def plot_convergence():
    ns = np.array([80, 160, 320, 640])
    ev_space = np.array([get_event(int(n), 7.5) for n in ns])
    dts = np.array([15.0, 7.5, 3.75])
    ev_time = np.array([get_event(320, float(d)) for d in dts])
    fig, ax = plt.subplots(1, 2, figsize=(6.7, 2.75), constrained_layout=True)
    ax[0].plot(ns, ev_space, marker="o", color=COLORS[0], lw=1.6)
    ax[0].set(xlabel="cells $N$", ylabel="event time / h", title="(a) Spatial refinement")
    ax[0].set_xscale("log", base=2)
    ax[1].plot(dts, ev_time, marker="s", color=COLORS[1], lw=1.6)
    ax[1].set(xlabel=r"time step $\Delta t$ / s", ylabel="event time / h", title="(b) Temporal refinement")
    ax[1].set_xscale("log", base=2)
    for a in ax:
        a.grid(alpha=0.25)
        a.spines[["top", "right"]].set_visible(False)
    savefig("fig05_convergence")


def plot_q4():
    wb = load_workbook(ATTACH / "附件2.xlsx", data_only=True, read_only=True)
    ws = wb.active
    vals = list(ws.iter_rows(min_row=2, values_only=True))
    tr = np.array([float(x[0]) / 3600 for x in vals])
    radius = np.array([float(x[1]) for x in vals])
    # Values sampled from the selected Q4 direct run N=160, dt=15 s.
    tq = np.array([0, 6, 12, 18, 24, 30, 36, 42, 48, 51.1028483])
    center = np.array([2.55, 1.7199, .7381, .4090, .2852, .2264, .1929, .1713, .1562, .1500])
    surface = np.array([2.55, .4208, .1670, .0892, .0674, .0595, .0560, .0541, .0530, .0526198])
    radius_event = np.interp(tq, tr, radius)
    fig, ax1 = plt.subplots(figsize=(6.7, 2.9), constrained_layout=True)
    ax1.plot(tr, radius, color=COLORS[2], lw=1.8, label="measured radius")
    ax1.set(xlabel="time / h", ylabel="radius / cm", title="Q4: shrinkage and moisture histories")
    ax1.set_xlim(0, 54)
    ax1.grid(alpha=0.25)
    ax1.spines["top"].set_visible(False)
    ax2 = ax1.twinx()
    ax2.plot(tq, center, color=COLORS[0], lw=1.7, label="center moisture")
    ax2.plot(tq, surface, color=COLORS[1], lw=1.5, ls="--", label="instantaneous surface")
    ax2.axhline(.15, color="#444444", lw=1, ls=":")
    ax2.set_ylabel(r"moisture $C$ / kg kg$^{-1}$")
    ax2.set_ylim(0, 2.75)
    lines = ax1.get_lines() + ax2.get_lines()
    ax1.legend(lines, [x.get_label() for x in lines], frameon=False, ncol=3, loc="upper right")
    savefig("fig06_q4_shrinkage")


def plot_compare():
    labels = ["Q3 fixed\nradius", "Q4 shrinking\nradius"]
    times = np.array([206914.89990234375, 183970.25390625]) / 3600
    errs = np.array([60.0, 70.0]) / 3600
    fig, ax = plt.subplots(figsize=(4.7, 2.8), constrained_layout=True)
    bars = ax.bar(labels, times, yerr=errs, capsize=4, color=[COLORS[0], COLORS[2]], width=.55,
                  error_kw=dict(ecolor="#444444", lw=1))
    ax.set(ylabel="first full-domain event time / h", title="Fixed-radius and Q4 candidate comparison")
    ax.grid(axis="y", alpha=0.25)
    ax.spines[["top", "right"]].set_visible(False)
    for b, y in zip(bars, times):
        ax.text(b.get_x() + b.get_width()/2, y + .8, f"{y:.2f}", ha="center", va="bottom", fontsize=9)
    savefig("fig07_q3_q4_compare")


def plot_sensitivity_events():
    """OAT ±20% event-time sensitivity for Q3 and Q4."""
    def load(path):
        return read_csv(HERE.parent / "sensitivity" / path)
    rows3 = load("q3_event_sensitivity.csv")
    rows4 = load("q4_event_sensitivity.csv")
    params = ["D", "h", "hm"]
    fig, ax = plt.subplots(1, 2, figsize=(6.7, 2.75), constrained_layout=True, sharey=True)
    for k, rows in enumerate((rows3, rows4)):
        base = next(f(r, "event_time_h") for r in rows if r["parameter"] == "D" and abs(f(r, "factor")-1) < 1e-8)
        for j, p in enumerate(params):
            vals = [f(r, "event_time_h") for r in rows if r["parameter"] == p]
            fac = [f(r, "factor") for r in rows if r["parameter"] == p]
            rel = [(v-base)/base*100 for v in vals]
            ax[k].plot(fac, rel, marker="o", color=COLORS[j], ls=STYLES[j], lw=1.55, label=p)
        ax[k].axhline(0, color="#555555", lw=.8)
        ax[k].set(xlabel="multiplicative factor", title=f"({'a' if k == 0 else 'b'}) {'Q3' if k == 0 else 'Q4'} event time")
        ax[k].set_xticks([.8, 1.0, 1.2])
        ax[k].grid(alpha=.25)
        ax[k].spines[["top", "right"]].set_visible(False)
    ax[0].set_ylabel("change from baseline / %")
    ax[1].legend(frameon=False, title="parameter", loc="upper left")
    savefig("fig08_sensitivity_events")


def plot_sensitivity_short():
    """Short-time Q2/Q3 field sensitivity at 1800 s."""
    rows = read_csv(HERE.parent / "sensitivity" / "short_q2q3_sensitivity.csv")
    params = ["D", "h", "hm"]
    baseline = {p: next(r for r in rows if r["parameter"] == p and abs(f(r, "factor")-1) < 1e-8) for p in params}
    fig, ax = plt.subplots(1, 2, figsize=(6.7, 2.75), constrained_layout=True)
    for j, p in enumerate(params):
        b = baseline[p]
        factors = [f(r, "factor") for r in rows if r["parameter"] == p]
        dT = [(f(r, "surface_temperature_C")-f(b, "surface_temperature_C")) for r in rows if r["parameter"] == p]
        dC = [(f(r, "surface_moisture")-f(b, "surface_moisture")) for r in rows if r["parameter"] == p]
        ax[0].plot(factors, dT, marker="o", color=COLORS[j], ls=STYLES[j], lw=1.55, label=p)
        ax[1].plot(factors, dC, marker="o", color=COLORS[j], ls=STYLES[j], lw=1.55, label=p)
    for a in ax:
        a.axhline(0, color="#555555", lw=.8)
        a.set_xticks([.8, 1.0, 1.2])
        a.grid(alpha=.25)
        a.spines[["top", "right"]].set_visible(False)
        a.set_xlabel("multiplicative factor")
    ax[0].set_ylabel(r"surface $\Delta T$ / $^\circ$C")
    ax[1].set_ylabel(r"surface $\Delta C$ / kg kg$^{-1}$")
    ax[0].set_title("(a) Temperature response at 1800 s")
    ax[1].set_title("(b) Moisture response at 1800 s")
    ax[0].legend(frameon=False, title="parameter")
    savefig("fig09_sensitivity_short")


if __name__ == "__main__":
    plot_environment()
    plot_q1()
    plot_q2()
    plot_q3()
    plot_convergence()
    plot_q4()
    plot_compare()
    plot_sensitivity_events()
    plot_sensitivity_short()
    print("generated", len(list(OUT.glob("fig*.pdf"))), "PDF figures in", OUT)
