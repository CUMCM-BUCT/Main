import csv
import hashlib
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from a_model.q2q3_xlsx import CSV_HEADER, CSV_NAMES, export_workbooks, export_workbooks_fast, inspect_sources


class Q2Q3XlsxTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.source = self.root / "source"
        self.templates = self.root / "templates"
        self.source.mkdir()
        self.templates.mkdir()
        result2 = Workbook()
        result2.active.title = "温度"
        result2.create_sheet("水分浓度")
        for sheet in result2.worksheets:
            sheet.append(["时间\\到药材中心的距离", 0])
            sheet.append([1, None])
            sheet.column_dimensions["A"].width = 19.625
            sheet.column_dimensions["B"].width = 9.625
        result2.save(self.templates / "result2.xlsx")
        result3 = Workbook()
        result3.active.title = "Sheet1"
        result3.active.append(["时间\\到药材中心的距离", 0])
        result3.active.append([60, None])
        result3.active.column_dimensions["A"].width = 19.625
        result3.active.column_dimensions["B"].width = 9.625
        result3.save(self.templates / "result3.xlsx")

    def tearDown(self):
        self.temporary.cleanup()

    def write_csvs(self):
        times = {
            CSV_NAMES[0]: (1., 2., 2.5),
            CSV_NAMES[1]: (1., 2., 2.5),
            CSV_NAMES[2]: (2.5,),
        }
        records = {}
        for index, name in enumerate(CSV_NAMES):
            path = self.source / name
            with path.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(CSV_HEADER)
                for time_s in times[name]:
                    writer.writerow((time_s, *(index + radius / 10 for radius in range(21))))
            records[name] = {
                "rows": len(times[name]),
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            }
        metadata = {
            "result_classification": "formal",
            "numerical_settings": {"convergence_verified": True},
            "source_files_unchanged": True,
            "result": {"status": "event", "event": {"state": {"time_s": 2.5}}},
            "checkpoint": {"files": records},
        }
        (self.source / "metadata.json").write_text(json.dumps(metadata), encoding="utf-8")

    def test_streaming_export_preserves_contract(self):
        self.write_csvs()
        output = self.root / "output"
        result2, result3 = export_workbooks(
            self.source, self.templates / "result2.xlsx", self.templates / "result3.xlsx", output
        )
        for path, sheets, rows in ((result2, ["温度", "水分浓度"], 4), (result3, ["Sheet1"], 2)):
            book = load_workbook(path, data_only=False, read_only=True)
            try:
                self.assertEqual(book.sheetnames, sheets)
                for sheet in book.worksheets:
                    values = list(sheet.iter_rows())
                    self.assertEqual(len(values), rows)
                    self.assertEqual(len(values[0]), 22)
                    self.assertEqual(values[0][0].value, "时间\\到药材中心的距离")
                    self.assertEqual(values[0][-1].value, 2)
                    self.assertEqual(values[1][1].number_format, "0.0000")
                    self.assertIn(values[1][1].data_type, ("n", "d"))
            finally:
                book.close()

    def test_candidate_metadata_is_rejected_without_outputs(self):
        self.write_csvs()
        path = self.source / "metadata.json"
        metadata = json.loads(path.read_text(encoding="utf-8"))
        metadata["result_classification"] = "candidate_only"
        path.write_text(json.dumps(metadata), encoding="utf-8")
        output = self.root / "output"
        with self.assertRaisesRegex(ValueError, "formal"):
            export_workbooks(self.source, self.templates / "result2.xlsx", self.templates / "result3.xlsx", output)
        self.assertFalse((output / "result2.xlsx").exists())

    def test_fast_streaming_export_preserves_contract(self):
        self.write_csvs()
        output = self.root / "fast-output"
        result2, result3 = export_workbooks_fast(
            self.source, self.templates / "result2.xlsx", self.templates / "result3.xlsx", output
        )
        self.assertTrue(result2.is_file())
        self.assertTrue(result3.is_file())
        result2_book = load_workbook(result2, read_only=True)
        result3_book = load_workbook(result3, read_only=True)
        try:
            self.assertEqual(result2_book.sheetnames, ["温度", "水分浓度"])
            self.assertEqual(result3_book.sheetnames, ["Sheet1"])
        finally:
            result2_book.close()
            result3_book.close()

    def test_off_cadence_row_must_be_terminal(self):
        self.write_csvs()
        path = self.source / CSV_NAMES[0]
        with path.open(encoding="utf-8") as stream:
            rows = list(csv.reader(stream))
        rows.append(["3", *("0" for _ in range(21))])
        with path.open("w", encoding="utf-8", newline="") as stream:
            csv.writer(stream).writerows(rows)
        with self.assertRaisesRegex(ValueError, "after its off-cadence terminal"):
            inspect_sources(self.source)


if __name__ == "__main__":
    unittest.main()
